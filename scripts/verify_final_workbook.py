"""
Verification script for Track_02_Submission_Workbook_FINAL.xlsx
"""
import openpyxl

def verify_final():
    wb_orig = openpyxl.load_workbook('Track_02_Submission_Workbook.xlsx', data_only=False)
    wb_final = openpyxl.load_workbook('Track_02_Submission_Workbook_FINAL.xlsx', data_only=False)

    print("=" * 70)
    print("      FINAL SUBMISSION WORKBOOK AUDIT & INTEGRITY REPORT")
    print("=" * 70)

    # 1. Evaluator Sheets Check
    evaluator_sheets = ['05_REVIEW_RESULT', '06_LIVE_VALIDATION', '07_BATCH_MODERATION']
    evaluator_diffs = []
    for sname in evaluator_sheets:
        ws_orig = wb_orig[sname]
        ws_final = wb_final[sname]
        max_r = max(ws_orig.max_row, ws_final.max_row)
        max_c = max(ws_orig.max_column, ws_final.max_column)
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                v_orig = ws_orig.cell(row=r, column=c).value
                v_fin = ws_final.cell(row=r, column=c).value
                if v_orig != v_fin:
                    evaluator_diffs.append((sname, openpyxl.utils.get_column_letter(c) + str(r), v_orig, v_fin))

    assert len(evaluator_diffs) == 0, f"Evaluator sheets diff detected: {evaluator_diffs}"
    print("[PASS] Evaluator sheets (05, 06, 07) are 100% UNMODIFIED and identical.")

    # 2. Check Candidate Details
    ws00 = wb_final['00_START']
    print("\n--- 00_START Details ---")
    print(f"  Candidate ID   [B5]: {ws00['B5'].value}")
    print(f"  Candidate Name [B6]: {ws00['B6'].value}")
    print(f"  Candidate Email[F6]: {ws00['F6'].value}")
    print(f"  Sub Date       [B7]: {ws00['B7'].value}")
    print(f"  Repo URL       [F7]: {ws00['F7'].value}")
    print(f"  AI Tool        [F11]: {ws00['F11'].value}")
    print(f"  Checklist G29  [G29]: {ws00['G29'].value}")

    assert ws00['B5'].value == "127014023", f"Expected ID 127014023, got {ws00['B5'].value}"
    assert ws00['F6'].value == "127014023@sastra.ac.in", f"Expected email 127014023@sastra.ac.in, got {ws00['F6'].value}"
    print("[PASS] Candidate ID and Registered Email successfully verified.")

    # 3. Check Deliverable 9
    ws01 = wb_final['01_DELIVERABLES']
    print("\n--- 01_DELIVERABLES Deliverable 9 ---")
    print(f"  Deliv 9 Status [C13]: {ws01['C13'].value}")
    print(f"  Deliv 9 Path   [D13]: {ws01['D13'].value}")
    print(f"  Deliv 9 Note   [E13]: {ws01['E13'].value}")

    assert ws01['C13'].value == "Pending — candidate must record and upload the video."
    assert ws01['D13'].value == "Pending — candidate must record and upload the video."
    print("[PASS] Demo video correctly marked as Pending — candidate must record and upload the video.")

    # 4. Check Submission Notes & Product Narrative
    print("\n--- Submission Notes & Performance Tier Declaration ---")
    print("  Submission Notes (01_DELIVERABLES!A25):")
    print("   ", ws01['A25'].value)
    print("  Product Recommendation (02_COMPONENTS!A26):")
    print("   ", wb_final['02_COMPONENTS']['A26'].value)

    print("\n" + "=" * 70)
    print("     ALL FINAL VERIFICATION CHECKS PASSED SUCCESSFULLY!")
    print("=" * 70)

if __name__ == '__main__':
    verify_final()
