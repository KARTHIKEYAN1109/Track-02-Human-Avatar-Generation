"""
Programmatic verification and audit of Track_02_Submission_Workbook_FILLED.xlsx
"""
import openpyxl

def verify():
    wb_orig = openpyxl.load_workbook('Track_02_Submission_Workbook.xlsx', data_only=False)
    wb_filled = openpyxl.load_workbook('Track_02_Submission_Workbook_FILLED.xlsx', data_only=False)

    print("=" * 70)
    print("           WORKBOOK INTEGRITY AND COMPLIANCE AUDIT")
    print("=" * 70)
    print("Filled workbook sheet list:", wb_filled.sheetnames)

    # 1. Check evaluator sheets (05, 06, 07)
    evaluator_sheets = ['05_REVIEW_RESULT', '06_LIVE_VALIDATION', '07_BATCH_MODERATION']
    evaluator_diffs = []

    for sname in evaluator_sheets:
        ws_orig = wb_orig[sname]
        ws_fill = wb_filled[sname]
        
        max_r = max(ws_orig.max_row, ws_fill.max_row)
        max_c = max(ws_orig.max_column, ws_fill.max_column)
        
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                v_orig = ws_orig.cell(row=r, column=c).value
                v_fill = ws_fill.cell(row=r, column=c).value
                if v_orig != v_fill:
                    evaluator_diffs.append((sname, openpyxl.utils.get_column_letter(c) + str(r), v_orig, v_fill))

    if evaluator_diffs:
        print(f"\n[FAIL] Evaluator sheet differences detected! Count: {len(evaluator_diffs)}")
        for d in evaluator_diffs:
            print("  ", d)
    else:
        print("\n[PASS] Evaluator sheets (05_REVIEW_RESULT, 06_LIVE_VALIDATION, 07_BATCH_MODERATION) are 100% UNTOUCHED and identical.")

    # 2. Inspect 00_START
    print("\n" + "-" * 70)
    print("SHEET 00_START AUDIT:")
    ws00 = wb_filled['00_START']
    start_fields = [
        ('Candidate ID', 'B5'),
        ('Candidate Name', 'B6'),
        ('Candidate Email', 'F6'),
        ('Submission Date', 'B7'),
        ('Repository URL', 'F7'),
        ('CPU / Cores', 'B8'),
        ('RAM (GB)', 'F8'),
        ('Operating System', 'B9'),
        ('Python / Runtime', 'F9'),
        ('Approved Free Compute Used', 'B10'),
        ('Provider / Runtime', 'F10'),
        ('Report / Demo Folder URL', 'B11'),
        ('AI Coding Tool Used', 'F11'),
        ('Candidate Signature', 'B34'),
        ('Declaration Date', 'F34'),
    ]
    for label, coord in start_fields:
        val = ws00[coord].value
        print(f"  {label:<30} [{coord}]: {val}")

    print("\nChecklist Items (Rows 23-30):")
    for r in range(23, 31):
        item = ws00[f"E{r}"].value
        status = ws00[f"F{r}"].value
        note = ws00[f"G{r}"].value
        print(f"  Row {r}: {item:<35} -> [{status}] {note}")

    # 3. Inspect 01_DELIVERABLES
    print("\n" + "-" * 70)
    print("SHEET 01_DELIVERABLES AUDIT:")
    ws01 = wb_filled['01_DELIVERABLES']
    for r in range(5, 14):
        num = ws01[f"A{r}"].value
        deliv = ws01[f"B{r}"].value
        status = ws01[f"C{r}"].value
        path_url = ws01[f"D{r}"].value
        print(f"  Deliv {num}: {deliv[:38]:<40} -> [{status}] {path_url}")

    print("\nReproduction Commands:")
    for r in range(16, 23):
        label = ws01[f"A{r}"].value
        cmd = ws01[f"B{r}"].value
        print(f"  {label:<25}: {cmd}")

    print("\nSubmission Notes:")
    print(" ", ws01["A25"].value)

    # 4. Inspect 02_COMPONENTS
    print("\n" + "-" * 70)
    print("SHEET 02_COMPONENTS AUDIT:")
    ws02 = wb_filled['02_COMPONENTS']
    comp_count = 0
    for r in range(5, 26):
        c_name = ws02[f"A{r}"].value
        if c_name:
            comp_count += 1
            ver = ws02[f"B{r}"].value
            role = ws02[f"C{r}"].value
            lic = ws02[f"E{r}"].value
            comm = ws02[f"G{r}"].value
            exec_env = ws02[f"H{r}"].value
            print(f"  {comp_count:2d}. {c_name:<25} | {ver:<22} | Lic: {lic:<22} | Comm: {comm:<10} | Exec: {exec_env}")
    print(f"Total Components Documented: {comp_count}")
    print("\nProduct Recommendation Narrative:")
    print(" ", ws02["A26"].value)

    # 5. Inspect 03_TEST_EVIDENCE
    print("\n" + "-" * 70)
    print("SHEET 03_TEST_EVIDENCE AUDIT:")
    ws03 = wb_filled['03_TEST_EVIDENCE']
    test_count = 0
    for r in range(7, 28):
        t_id = ws03[f"A{r}"].value
        if t_id:
            test_count += 1
            lvl = ws03[f"B{r}"].value
            expected = ws03[f"D{r}"].value
            route = ws03[f"G{r}"].value
            passed = ws03[f"H{r}"].value
            m1 = ws03[f"I{r}"].value
            v1 = ws03[f"J{r}"].value
            rt = ws03[f"M{r}"].value
            print(f"  {test_count:2d}. {t_id:<18} [{lvl:<8}] Pass: {passed:<5} | {m1}: {v1:<10} | Time: {rt}s | Route: {route}")
    print(f"Total Test Executions Documented: {test_count}")
    print("\nBenchmark Method Narrative:")
    print(" ", ws03["A28"].value)

    # 6. Inspect 04_TRACK_CRITERIA
    print("\n" + "-" * 70)
    print("SHEET 04_TRACK_CRITERIA AUDIT:")
    ws04 = wb_filled['04_TRACK_CRITERIA']
    print("  Baseline requirement (B12):", ws04["B12"].value[:70], "...")
    print("  Strong requirement (B13):  ", ws04["B13"].value[:70], "...")
    print("  Exceptional requirement (B14):", ws04["B14"].value[:70], "...")

    print("\n" + "=" * 70)
    print("                ALL AUDIT CHECKS COMPLETED")
    print("=" * 70)

if __name__ == '__main__':
    verify()
